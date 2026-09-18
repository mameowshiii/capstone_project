from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("Barangay_Pili_Questionnaires_Portrait_Format.docx")

ISO_SECTIONS = [
    ("Functional Suitability", [
        "The system provides the needed functions for requesting barangay clearances and certificates.",
        "The system correctly records resident information and document requests.",
        "The request-tracking feature provides accurate request statuses.",
        "The system allows administrators to manage residents, requests, certificates, announcements, and records effectively.",
        "The system supports the required barangay processes completely.",
    ]),
    ("Performance Efficiency", [
        "Pages and system features load within an acceptable time.",
        "The system responds quickly when information is submitted or updated.",
        "Searching resident records and requests is fast and efficient.",
        "The system remains responsive while handling multiple transactions.",
        "The system reduces the time needed to process barangay document requests.",
    ]),
    ("Compatibility", [
        "The system works properly using commonly available web browsers.",
        "The system can be accessed and used on mobile devices.",
        "The system displays information properly across different screen sizes.",
        "The web system and resident mobile application provide consistent information.",
        "The system works well with the barangay's available devices and internet connection.",
    ]),
    ("Usability", [
        "The system's interface is easy to understand.",
        "Menus, buttons, and labels are clear and easy to locate.",
        "I can complete tasks in the system with minimal assistance.",
        "Error messages and instructions are understandable.",
        "The system is visually organized and pleasant to use.",
    ]),
    ("Reliability", [
        "The system performs its functions consistently without unexpected errors.",
        "Submitted requests and resident records are saved correctly.",
        "The system provides accurate information after updates are made.",
        "The system can recover properly from minor errors or interrupted actions.",
        "The system is dependable for daily barangay transactions.",
    ]),
    ("Security", [
        "The system requires authorized users to log in before accessing protected information.",
        "Resident information is accessible only to authorized users.",
        "User accounts and passwords are adequately protected.",
        "The system maintains records of important administrative activities.",
        "The system protects sensitive resident and transaction data.",
    ]),
    ("Maintainability (IT evaluators and system administrators only)", [
        "The system is organized in a way that makes updates manageable.",
        "System errors can be identified and corrected efficiently.",
        "New features or modules can be added without greatly affecting existing functions.",
        "The system is easy to test after changes are made.",
        "The system documentation and structure support future maintenance.",
    ]),
    ("Portability", [
        "The system can be used on different devices with minimal changes.",
        "The system can be deployed in another similar barangay environment.",
        "The system can be accessed through different supported platforms.",
        "The system can adapt to future hardware or software upgrades.",
        "The system can be installed and configured with reasonable effort.",
    ]),
]

USE_SECTIONS = [
    ("Usefulness", [
        "The system helps me complete barangay-related tasks more quickly.",
        "The system improves the way document requests are processed.",
        "The system makes it easier to monitor the status of requests.",
        "The system helps reduce manual paperwork and repetitive transactions.",
        "The system makes resident and request information easier to manage.",
        "The system is useful for accessing barangay announcements and updates.",
        "The system improves communication between residents and barangay staff.",
        "Overall, the system meets my needs.",
    ]),
    ("Ease of Use", [
        "The system is simple to use.",
        "I can use the system without difficulty.",
        "The features work in the way I expect them to.",
        "It is easy to find the information or feature I need.",
        "Completing a request is straightforward.",
        "The system's navigation is clear.",
        "I can correct mistakes easily when entering information.",
        "The system requires only a few steps to complete common tasks.",
    ]),
    ("Ease of Learning", [
        "I learned how to use the system quickly.",
        "It was easy to remember how to use the system after my first use.",
        "The system's instructions helped me learn its features.",
        "I could become productive with the system in a short time.",
        "New users can learn to use the system easily.",
        "The system does not require extensive training.",
    ]),
    ("Satisfaction", [
        "I am satisfied with the system.",
        "I would recommend the system to other residents or barangay staff.",
        "I feel confident when using the system.",
        "The system is pleasant to use.",
        "The system makes barangay transactions more convenient.",
        "I would like to continue using the system.",
        "The system meets my expectations.",
        "Overall, I am satisfied with the Barangay Pili Clearance and Certificate Management System.",
    ]),
]

CUSTOM_SECTIONS = [
    ("Access and Account Experience", [
        "I can access the system when I need barangay services.",
        "Account registration and email verification are clear and manageable.",
        "Logging in and recovering a forgotten password are easy to complete.",
        "My resident profile information is easy to review and update.",
    ]),
    ("Clearance and Certificate Requests", [
        "The system clearly explains the requirements for each clearance or certificate.",
        "The online request form asks only for information needed to process my transaction.",
        "Uploading supporting documents and proof of payment is straightforward.",
        "I can review the details of a request before and after submitting it.",
        "The system makes requesting and releasing barangay documents more organized.",
    ]),
    ("Tracking and Notifications", [
        "The request status labels are clear and easy to understand.",
        "The system provides timely updates when the status of my request changes.",
        "Email or SMS notifications contain enough information about the next step I should take.",
        "The tracking feature reduces the need to visit or contact the barangay office for updates.",
    ]),
    ("Other Barangay Services and Information", [
        "Barangay announcements and bulletins are timely and easy to understand.",
        "The borrowing feature clearly shows the item, quantity, schedule, and request status.",
        "The summons feature presents hearing schedules and case-related updates clearly.",
        "The system makes important barangay information easier for residents to access.",
    ]),
    ("Privacy, Trust, and Overall Acceptance", [
        "I trust the system to handle my personal and transaction information appropriately.",
        "The system gives me confidence that my submitted information reaches the barangay office.",
        "The system reduces the time and effort required for barangay transactions.",
        "I prefer using this system over a purely manual request process.",
        "I would use this system for future barangay transactions.",
        "The system is beneficial to Barangay Pili residents and staff.",
    ]),
]

OPEN_ENDED = [
    "What is the main strength of the system?",
    "What difficulty or problem did you encounter while using the system?",
    "Which feature or process should be improved first? Please explain your recommendation.",
    "What additional feature or barangay service should be added?",
    "Other feedback or suggestions:",
]

RECOMMENDATION_OPTIONS = [
    "I recommend the system for implementation or continued use.",
    "I recommend the system after minor improvements.",
    "I recommend further development and testing before implementation.",
    "I do not recommend the system at this time.",
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def style_run(run, size=10, bold=False, color=None, font="Times New Roman"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7 if level == 1 else 4)
    p.paragraph_format.space_after = Pt(3)
    keep_with_next(p)
    run = p.add_run(text)
    if level == 1:
        style_run(run, size=12, bold=True, color=(31, 78, 54))
    else:
        style_run(run, size=10.5, bold=True, color=(31, 78, 54))
    return p


def add_text(doc, text, bold_label=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    if bold_label and text.startswith(bold_label):
        r1 = p.add_run(bold_label)
        style_run(r1, size=10, bold=True)
        r2 = p.add_run(text[len(bold_label):])
        style_run(r2, size=10)
    else:
        style_run(p.add_run(text), size=10)
    return p


def style_table_borders(table, color="7F8C8D", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:color"), color)


def add_rating_table(doc, items, start_number):
    headers = ["No.", "Statement", "5", "4", "3", "2", "1", "N/A"]
    widths = [0.38, 4.42, 0.42, 0.42, 0.42, 0.42, 0.42, 0.52]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    style_table_borders(table)

    header = table.rows[0]
    repeat_table_header(header)
    for i, (cell, label, width) in enumerate(zip(header.cells, headers, widths)):
        set_cell_width(cell, width)
        set_cell_shading(cell, "1F4E36")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        style_run(p.add_run(label), size=9, bold=True, color=(255, 255, 255))

    for offset, statement in enumerate(items):
        row = table.add_row()
        prevent_row_split(row)
        values = [str(start_number + offset), statement, "☐", "☐", "☐", "☐", "☐", "☐"]
        for i, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if offset % 2:
                set_cell_shading(cell, "EDF3EF")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 1 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            style_run(p.add_run(value), size=9, font="Arial" if i >= 2 else "Times New Roman")

    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return start_number + len(items)


def add_questionnaire_section(doc, title, intro, sections):
    doc.add_page_break()
    add_heading(doc, title, level=1)
    add_text(doc, intro)
    number = 1
    for subsection, items in sections:
        add_heading(doc, subsection, level=2)
        number = add_rating_table(doc, items, number)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    style_run(p.add_run("BARANGAY PILI CLEARANCE AND CERTIFICATE MANAGEMENT SYSTEM"), 14, True, (31, 78, 54))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    style_run(p.add_run("SYSTEM EVALUATION QUESTIONNAIRES"), 13, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    style_run(p.add_run("Barangay Pili, Madridejos, Cebu"), 10, False)

    add_heading(doc, "Purpose", level=1)
    add_text(doc, "This instrument gathers feedback on the quality, usability, and acceptance of the Barangay Pili Clearance and Certificate Management System. Participation is voluntary. Responses will be kept confidential and used only for academic system evaluation.")

    add_heading(doc, "Respondent Profile", level=1)
    profile = doc.add_table(rows=2, cols=4)
    profile.alignment = WD_TABLE_ALIGNMENT.CENTER
    profile.autofit = False
    style_table_borders(profile)
    profile_values = [
        ["Respondent type", "☐ Resident/User   ☐ Barangay Staff/Admin   ☐ IT Expert", "Date", "________________"],
        ["Age", "__________", "Sex", "________________"],
    ]
    profile_widths = [1.05, 4.0, 0.6, 1.75]
    for row, values in zip(profile.rows, profile_values):
        for index, (cell, value, width) in enumerate(zip(row.cells, values, profile_widths)):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index in (0, 2):
                set_cell_shading(cell, "D9E8DE")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            style_run(p.add_run(value), 9.5, index in (0, 2))

    add_heading(doc, "Rating Scale", level=1)
    scale = doc.add_table(rows=2, cols=6)
    scale.alignment = WD_TABLE_ALIGNMENT.CENTER
    scale.autofit = False
    style_table_borders(scale)
    scale_values = [
        ["5", "4", "3", "2", "1", "N/A"],
        ["Strongly Agree", "Agree", "Neutral", "Disagree", "Strongly Disagree", "Not Applicable"],
    ]
    for row_index, row in enumerate(scale.rows):
        for cell, value in zip(row.cells, scale_values[row_index]):
            set_cell_width(cell, 1.23)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, "1F4E36")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            style_run(p.add_run(value), 9, True, (255, 255, 255) if row_index == 0 else None)

    add_text(doc, "Instructions: Put one check mark in the rating column that best represents your assessment. Select N/A only when you did not use or observe the feature described.", bold_label="Instructions:")

    add_questionnaire_section(
        doc,
        "A. ISO/IEC 25010-BASED SYSTEM EVALUATION QUESTIONNAIRE",
        "Rate the quality of the system using the statements below.",
        ISO_SECTIONS,
    )
    add_questionnaire_section(
        doc,
        "B. USE QUESTIONNAIRE",
        "Rate the system's usefulness, ease of use, ease of learning, and user satisfaction.",
        USE_SECTIONS,
    )
    add_questionnaire_section(
        doc,
        "C. SELF-MADE QUESTIONNAIRE: SYSTEM-SPECIFIC USER ACCEPTANCE AND FEEDBACK",
        "This researcher-developed questionnaire is tailored to the actual services and features of the Barangay Pili system. Select N/A for a module you did not use, such as borrowing or summons.",
        CUSTOM_SECTIONS,
    )

    add_heading(doc, "Researcher's Note", level=2)
    add_text(doc, "Section A is organized around ISO/IEC 25010 software-quality characteristics. Section B uses the USE dimensions of usefulness, ease of use, ease of learning, and satisfaction. Section C is a self-made instrument designed specifically for the Barangay Pili Clearance and Certificate Management System. The instrument should be reviewed and pilot-tested before formal data collection.")

    doc.add_page_break()
    add_heading(doc, "D. RECOMMENDATION AND FEEDBACK", level=1)
    add_text(doc, "Please provide your overall recommendation and specific feedback about the system. Your responses will guide future improvements.")

    add_heading(doc, "Overall Recommendation", level=2)
    recommendation = doc.add_table(rows=1, cols=2)
    recommendation.alignment = WD_TABLE_ALIGNMENT.CENTER
    recommendation.autofit = False
    style_table_borders(recommendation)
    for cell, label, width in zip(recommendation.rows[0].cells, ["Select one", "Recommendation"], [1.0, 6.4]):
        set_cell_width(cell, width)
        set_cell_shading(cell, "1F4E36")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        style_run(p.add_run(label), 9, True, (255, 255, 255))
    for index, option in enumerate(RECOMMENDATION_OPTIONS):
        row = recommendation.add_row()
        prevent_row_split(row)
        for cell, width in zip(row.cells, [1.0, 6.4]):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = row.cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        style_run(p.add_run("☐"), 10, font="Arial")
        p = row.cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        style_run(p.add_run(option), 9)
        if index % 2:
            set_cell_shading(row.cells[0], "EDF3EF")
            set_cell_shading(row.cells[1], "EDF3EF")

    add_heading(doc, "Written Feedback and Recommendations", level=2)
    feedback = doc.add_table(rows=1, cols=2)
    feedback.alignment = WD_TABLE_ALIGNMENT.CENTER
    feedback.autofit = False
    style_table_borders(feedback)
    for cell, label, width in zip(feedback.rows[0].cells, ["Guide Question", "Response"], [3.2, 4.2]):
        set_cell_width(cell, width)
        set_cell_shading(cell, "1F4E36")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        style_run(p.add_run(label), 9, True, (255, 255, 255))
    repeat_table_header(feedback.rows[0])
    for index, question in enumerate(OPEN_ENDED, 1):
        row = feedback.add_row()
        prevent_row_split(row)
        for cell, width in zip(row.cells, [3.2, 4.2]):
            set_cell_width(cell, width)
        p = row.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        style_run(p.add_run(f"{index}. {question}"), 9)
        p = row.cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(34)
        style_run(p.add_run(""), 9)
        if index % 2 == 0:
            set_cell_shading(row.cells[0], "EDF3EF")
            set_cell_shading(row.cells[1], "EDF3EF")

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(footer.add_run("Barangay Pili System Evaluation Questionnaire"), 8, False, (90, 90, 90))

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(f"Created: {OUTPUT}")
